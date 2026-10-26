from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from knowledge_mcp.ingest import (
    chunk_segments,
    discover_files,
    parse_docx,
    parse_markdown,
    parse_text,
    processing_fingerprint,
)
from knowledge_mcp.models import Segment


class IngestTests(unittest.TestCase):
    def test_discovery_skips_hidden_dependencies_and_symlinks(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "docs").mkdir()
            (root / "docs" / "keep.md").write_text("keep", encoding="utf-8")
            (root / ".git").mkdir()
            (root / ".git" / "hidden.md").write_text("hidden", encoding="utf-8")
            (root / "node_modules").mkdir()
            (root / "node_modules" / "vendor.txt").write_text("vendor", encoding="utf-8")
            (root / "linked.md").symlink_to(root / "docs" / "keep.md")

            result = discover_files(
                root,
                {".md": "markdown", ".txt": "text"},
                ("**/node_modules/**",),
            )
            self.assertTrue(result.complete)
            self.assertEqual([item.relative_path for item in result.files], ["docs/keep.md"])

    def test_markdown_chunks_fit_token_budget_and_keep_locations(self) -> None:
        content = b"# Heading\n\n" + (b"word " * 30) + b"\n"
        segments = parse_markdown(content, 10)
        chunks = chunk_segments(
            segments,
            "notes/test.md",
            "md",
            "hash",
            "process",
            lambda text: len(text.split()) + 2,
            max_tokens=10,
            overlap_tokens=2,
        )
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.text.split()) + 2 <= 10 for chunk in chunks))
        self.assertTrue(all(chunk.metadata["path"] == "notes/test.md" for chunk in chunks))


class StructuralChunkingTests(unittest.TestCase):
    def chunks(self, source, strategy="paragraph", budget=30, overlap=0, counter=None):
        segments = parse_text(source.encode("utf-8"), 10) if isinstance(source, str) else source
        return chunk_segments(
            segments,
            "docs/example.txt",
            "txt",
            "hash",
            "fingerprint",
            counter or (lambda text: len(text.split()) + 2),
            budget,
            overlap,
            strategy,
        )

    def test_modes_preserve_paragraphs_sections_or_pack(self):
        text = "# Heading\n\nFirst paragraph.\n\nSecond paragraph.\n\n# Next\n\nThird paragraph."
        paragraph = self.chunks(text)
        section = self.chunks(text, "section")
        packed = self.chunks(text, "packed")
        self.assertEqual(len(paragraph), 3)
        self.assertEqual(len(section), 2)
        self.assertEqual(len(packed), 1)
        self.assertEqual("".join(c.text for c in paragraph), text)
        self.assertEqual(paragraph[0].text.count("# Heading"), 1)
        self.assertNotIn("# Heading", paragraph[1].text)

    def test_section_falls_back_to_paragraphs_without_headings(self):
        self.assertEqual(len(self.chunks("First.\n\nSecond.", "section")), 2)

    def test_large_section_falls_back_to_independent_paragraphs(self):
        text = "# Title\n\nOne two three.\n\nFour five six."
        chunks = self.chunks(text, "section", budget=7)
        self.assertEqual(len(chunks), 2)
        self.assertEqual("".join(c.text for c in chunks), text)

    def test_long_paragraph_groups_whole_sentences_without_overlap(self):
        text = "One two three. Four five six. Seven eight nine."
        chunks = self.chunks(text, budget=8, overlap=4)
        self.assertEqual(
            [c.text.strip() for c in chunks], ["One two three. Four five six.", "Seven eight nine."]
        )
        self.assertEqual("".join(c.text for c in chunks), text)

    def test_lists_remain_intact_and_split_by_items_when_large(self):
        text = "Intro.\n\n- First item\n  continuation\n- Second item\n- Third item\n"
        self.assertEqual(len(self.chunks(text)), 2)
        chunks = self.chunks(text, budget=8, overlap=4)
        self.assertEqual(len(chunks), 3)
        self.assertIn("continuation", chunks[1].text)
        self.assertEqual("".join(c.text for c in chunks), text)

    def test_numbered_lists_are_not_split_as_sentences(self):
        text = "1. First item\n2) Second item\n3. Third item"
        chunks = self.chunks(text, budget=8)
        self.assertEqual(len(chunks), 2)
        self.assertEqual("".join(c.text for c in chunks), text)
        self.assertTrue(chunks[1].text.startswith("3."))

    def test_fences_ignore_blank_lines_headings_and_sentence_punctuation(self):
        text = "```python\na = 1.0\n\n# not a heading\nb = 2.0\n```\n\nEnd."
        segments = parse_markdown(text.encode(), 10)
        self.assertEqual([s.kind for s in segments], ["code", "paragraph"])
        chunks = self.chunks(segments, budget=8)
        self.assertEqual("".join(c.text for c in chunks), text)
        self.assertTrue(all(len(c.text.split()) + 2 <= 8 for c in chunks))

    def test_fence_marker_with_trailing_text_is_not_a_closing_fence(self):
        text = "```text\n```not-closing\n\n# still code\n```\n\nAfter."
        segments = parse_markdown(text.encode(), 10)
        self.assertEqual([s.kind for s in segments], ["code", "paragraph"])
        self.assertIn("# still code", segments[0].text)
        self.assertEqual("".join(s.text for s in segments), text)

    def test_pdf_page_boundaries_are_preserved_in_every_mode(self):
        segments = [Segment("First.\n\nSecond.", page=1), Segment("Third.", page=2)]
        for strategy in ("paragraph", "section", "packed"):
            with self.subTest(strategy=strategy):
                chunks = self.chunks(segments, strategy)
                self.assertEqual(chunks[-1].metadata["page"], 2)
                self.assertTrue(all(c.metadata["page"] == 1 for c in chunks[:-1]))
                self.assertEqual(chunks[-1].text, "Third.")

    def test_line_locations_crlf_and_unicode_are_preserved(self):
        text = "# Título\r\n\r\nOlá mundo.\r\n\r\nÚltimo parágrafo."
        chunks = self.chunks(text)
        self.assertEqual("".join(c.text for c in chunks), text)
        self.assertEqual(chunks[0].metadata["start_line"], 1)
        self.assertEqual(chunks[0].metadata["end_line"], 3)
        self.assertEqual(chunks[1].metadata["start_line"], 5)
        split = self.chunks("aaa\nbbb\nccc", budget=6, counter=lambda t: len(t) + 2)
        self.assertEqual([c.metadata["start_line"] for c in split], [1, 2, 3])
        self.assertEqual([c.metadata["end_line"] for c in split], [1, 2, 3])

    def test_forced_splits_preserve_characters_with_zero_overlap(self):
        for text in ("abcdefghijklmno", "áβ中🙂" * 8, "aa bb cc dd ee", "x" * 100):
            with self.subTest(text=text):
                chunks = self.chunks(text, budget=7, counter=lambda t: len(t) + 2)
                self.assertEqual("".join(c.text for c in chunks), text)
                self.assertTrue(all(len(c.text) + 2 <= 7 for c in chunks))

    def test_forced_overlap_is_bounded_and_makes_progress(self):
        text = "abcdefghijklmnop"
        chunks = self.chunks(text, budget=7, overlap=4, counter=lambda t: len(t) + 2)
        self.assertEqual(chunks[0].text, "abcde")
        self.assertEqual(chunks[1].text, "defgh")
        self.assertTrue(chunks[-1].text.endswith("p"))
        self.assertLess(len(chunks), len(text))
        self.assertTrue(all(len(c.text) + 2 <= 7 for c in chunks))

    def test_impossible_budget_errors_instead_of_emitting_oversized_chunk(self):
        with self.assertRaisesRegex(ValueError, "Token budget"):
            self.chunks("abc", budget=2, counter=lambda t: len(t) + 2)

    def test_exact_budget_and_empty_input(self):
        self.assertEqual(len(self.chunks("one two three", budget=5)), 1)
        for text in ("", " \n\n\t"):
            self.assertEqual(self.chunks(text), [])

    def test_fingerprint_includes_strategy_and_chunker_version(self):
        from unittest.mock import patch

        first = processing_fingerprint("text", 256, 40, "paragraph")
        self.assertNotEqual(first, processing_fingerprint("text", 256, 40, "packed"))
        with patch("knowledge_mcp.ingest.CHUNKER_VERSION", 999):
            self.assertNotEqual(first, processing_fingerprint("text", 256, 40, "paragraph"))


class DocxIngestTests(unittest.TestCase):
    def content(self, document):
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def test_docx_body_lists_and_nested_tables_in_order(self):
        from docx import Document

        document = Document()
        document.add_heading("Title", level=1)
        document.add_paragraph("First paragraph.")
        document.add_paragraph("Item one", style="List Bullet")
        document.add_paragraph("Item two", style="List Bullet")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Cell A"
        table.cell(0, 1).text = "Cell B"
        nested = table.cell(0, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested"
        document.add_paragraph("Last paragraph.")
        document.sections[0].header.paragraphs[0].text = "Excluded header"
        segments = parse_docx(self.content(document), 10)
        self.assertEqual(
            [s.kind for s in segments], ["heading", "paragraph", "list", "table_row", "paragraph"]
        )
        self.assertEqual(segments[0].heading_level, 1)
        text = "".join(s.text for s in segments)
        self.assertLess(text.index("Item one"), text.index("Cell A"))
        self.assertLess(text.index("Nested"), text.index("Last paragraph"))
        self.assertNotIn("Excluded header", text)
        chunks = StructuralChunkingTests().chunks(segments)
        self.assertTrue(
            all("page" not in c.metadata and "start_line" not in c.metadata for c in chunks)
        )

    def test_merged_cells_are_not_duplicated(self):
        from docx import Document

        for rows, cols in ((1, 2), (2, 1)):
            with self.subTest(rows=rows, cols=cols):
                document = Document()
                table = document.add_table(rows=rows, cols=cols)
                table.cell(0, 0).merge(table.cell(rows - 1, cols - 1)).text = "Merged"
                text = "".join(s.text for s in parse_docx(self.content(document), 10))
                self.assertEqual(text.count("Merged"), 1)

    def test_empty_and_malformed_docx(self):
        from docx import Document

        self.assertEqual(parse_docx(self.content(Document()), 10), [])
        with self.assertRaisesRegex(ValueError, "DOCX could not be read"):
            parse_docx(b"not a zip archive", 10)


if __name__ == "__main__":
    unittest.main()
