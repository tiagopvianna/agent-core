from __future__ import annotations

import math
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from typing import Any

from knowledge_mcp.config import AppConfig, EmbeddingConfig, IndexConfig, ServerConfig
from knowledge_mcp.core import KnowledgeCore
from knowledge_mcp.models import Chunk, SearchHit


class FakeEmbedder:
    max_tokens = 32

    def __init__(self) -> None:
        self.document_calls = 0

    def count_tokens(self, text: str) -> int:
        return max(1, len(text.split())) + 2

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.document_calls += 1
        return [self._vector(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vector(text)

    def identity(self) -> dict[str, Any]:
        return {
            "provider": "fake",
            "model": "test",
            "dimension": 2,
            "normalize": True,
            "max_tokens": self.max_tokens,
        }

    @staticmethod
    def _vector(text: str) -> list[float]:
        lowered = text.lower()
        vector = [float(lowered.count("alpha") + 1), float(lowered.count("beta") + 1)]
        length = math.sqrt(sum(value * value for value in vector))
        return [value / length for value in vector]


class MemoryStore:
    def __init__(self) -> None:
        self.files: dict[str, dict[str, Any]] = {}
        self.documents: dict[str, tuple[list[Chunk], list[list[float]]]] = {}
        self.fail_replace = False
        self.last_sync: dict[str, Any] | None = None

    def initialize(self, _root: Path, _identity: dict[str, Any]) -> None:
        return None

    def file_entries(self) -> dict[str, dict[str, Any]]:
        return self.files

    def update_file(self, path: str, entry: dict[str, Any]) -> None:
        self.files[path] = entry

    def remove_file_entry(self, path: str) -> None:
        self.files.pop(path, None)

    def replace_document(
        self,
        path: str,
        chunks: list[Chunk],
        embeddings: list[list[float]],
        _batch_size: int,
    ) -> None:
        self.documents.pop(path, None)
        if self.fail_replace:
            raise RuntimeError("simulated write failure")
        self.documents[path] = (chunks, embeddings)

    def delete_document(self, path: str) -> None:
        self.documents.pop(path, None)

    def record_sync(self, report: dict[str, Any], _complete_scan: bool) -> None:
        self.last_sync = report

    def search(self, query: list[float], limit: int, folder: str | None = None) -> list[SearchHit]:
        hits: list[SearchHit] = []
        for path, (chunks, embeddings) in self.documents.items():
            if self.files.get(path, {}).get("state") != "active":
                continue
            if folder and not path.startswith(f"{folder.strip('/')}/"):
                continue
            for chunk, embedding in zip(chunks, embeddings, strict=True):
                similarity = sum(left * right for left, right in zip(query, embedding, strict=True))
                hits.append(
                    SearchHit(
                        path=path,
                        file_type=str(chunk.metadata["file_type"]),
                        snippet=chunk.text,
                        distance=1 - similarity,
                        metric="cosine",
                        indexed_at=self.files[path].get("indexed_at"),
                        chunk_id=chunk.id,
                    )
                )
        return sorted(hits, key=lambda hit: hit.distance)[:limit]

    def inspect(self, path: str, max_chars: int) -> dict[str, Any]:
        if self.files.get(path, {}).get("state") != "active":
            raise FileNotFoundError(path)
        chunks = self.documents[path][0]
        return {
            "path": path,
            "chunks": [{"text": chunk.text[:max_chars]} for chunk in chunks],
        }

    def status(self) -> dict[str, Any]:
        return {
            "indexed_files": sum(entry.get("state") == "active" for entry in self.files.values()),
            "indexed_chunks": sum(len(chunks) for chunks, _ in self.documents.values()),
            "last_sync": self.last_sync,
        }


class CoreTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        base = Path(self.temporary.name)
        self.source = base / "source"
        self.source.mkdir()
        self.config = AppConfig(
            source_root=self.source,
            data_dir=base / "data",
            server=ServerConfig(token_file=base / "token"),
            index=IndexConfig(
                max_files_per_sync=10,
                max_file_bytes=100_000,
                max_pdf_pages=50,
                chunk_tokens=24,
                overlap_tokens=4,
            ),
            embedding=EmbeddingConfig(provider="fake", model="test", batch_size=4),
        )
        self.embedder = FakeEmbedder()
        self.store = MemoryStore()
        self.core = KnowledgeCore(self.config, self.embedder, self.store)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_indexes_updates_and_removes_file(self) -> None:
        document = self.source / "apis" / "payments.md"
        document.parent.mkdir()
        document.write_text("# Payments\n\nAlpha retries are idempotent.\n", encoding="utf-8")

        first = self.core.sync()
        self.assertEqual(first["indexed"], 1)
        self.assertEqual(self.store.files["apis/payments.md"]["state"], "active")
        first_call_count = self.embedder.document_calls

        unchanged = self.core.sync()
        self.assertEqual(unchanged["indexed"], 0)
        self.assertEqual(self.embedder.document_calls, first_call_count)

        document.write_text("# Payments\n\nBeta retries use a token.\n", encoding="utf-8")
        changed = self.core.sync()
        self.assertEqual(changed["indexed"], 1)
        self.assertEqual(self.core.search("beta", folder="apis")[0]["path"], "apis/payments.md")

        document.unlink()
        removed = self.core.sync()
        self.assertEqual(removed["removed"], 1)
        self.assertNotIn("apis/payments.md", self.store.files)

    def test_interrupted_replacement_is_hidden_and_recovers(self) -> None:
        document = self.source / "note.txt"
        document.write_text("alpha", encoding="utf-8")
        self.assertEqual(self.core.sync()["indexed"], 1)

        document.write_text("beta", encoding="utf-8")
        self.store.fail_replace = True
        failed = self.core.sync()
        self.assertEqual(failed["failed"], 1)
        self.assertEqual(self.store.files["note.txt"]["state"], "pending_replace")
        self.assertEqual(self.core.search("beta"), [])

        self.store.fail_replace = False
        recovered = self.core.sync()
        self.assertEqual(recovered["indexed"], 1)
        self.assertEqual(self.store.files["note.txt"]["state"], "active")

    def test_missing_root_never_deletes_indexed_files(self) -> None:
        document = self.source / "note.txt"
        document.write_text("alpha", encoding="utf-8")
        self.assertEqual(self.core.sync()["indexed"], 1)

        self.source.rename(self.source.with_name("source-offline"))
        report = self.core.sync()
        self.assertFalse(report["scan_complete"])
        self.assertIn("note.txt", self.store.files)
        self.assertIn("note.txt", self.store.documents)

    def test_failed_content_is_not_retried_until_changed_or_poll_retry(self) -> None:
        document = self.source / "binary.txt"
        document.write_bytes(b"\xff")
        failed = self.core.sync()
        self.assertEqual(failed["failed"], 1)
        calls = self.embedder.document_calls

        skipped = self.core.sync(retry_errors=False)
        self.assertEqual(skipped["failed"], 0)
        self.assertEqual(self.embedder.document_calls, calls)

        retried = self.core.sync(retry_errors=True)
        self.assertEqual(retried["failed"], 1)

        document.write_text("alpha", encoding="utf-8")
        fixed = self.core.sync()
        self.assertEqual(fixed["indexed"], 1)

    def test_strategy_change_reindexes_unchanged_source(self) -> None:
        document = self.source / "note.txt"
        document.write_text("alpha\n\nbeta", encoding="utf-8")
        self.assertEqual(self.core.sync()["indexed"], 1)
        self.assertEqual(len(self.store.documents["note.txt"][0]), 2)
        config = replace(self.config, index=replace(self.config.index, chunk_strategy="packed"))
        core = KnowledgeCore(config, self.embedder, self.store)
        self.assertEqual(core.sync()["indexed"], 1)
        self.assertEqual(len(self.store.documents["note.txt"][0]), 1)
        self.assertEqual(core.sync()["indexed"], 0)

    def test_model_capacity_caps_configured_chunk_budget(self) -> None:
        document = self.source / "note.txt"
        document.write_text("alpha " * 100, encoding="utf-8")
        config = replace(self.config, index=replace(self.config.index, chunk_tokens=256))
        core = KnowledgeCore(config, self.embedder, self.store)
        self.assertEqual(core.sync()["indexed"], 1)
        self.assertTrue(
            all(
                self.embedder.count_tokens(c.text) <= self.embedder.max_tokens
                for c in self.store.documents["note.txt"][0]
            )
        )

    def test_docx_uses_existing_indexing_pipeline_and_file_size_limit(self) -> None:
        from docx import Document

        document = Document()
        document.add_heading("Alpha", level=1)
        document.add_paragraph("Beta is searchable.")
        document.save(self.source / "note.docx")
        self.assertEqual(self.core.sync()["indexed"], 1)
        self.assertEqual(self.core.search("beta")[0]["path"], "note.docx")
        config = replace(self.config, index=replace(self.config.index, max_file_bytes=1))
        (self.source / "other.docx").write_bytes((self.source / "note.docx").read_bytes())
        core = KnowledgeCore(config, self.embedder, self.store)
        report = core.sync()
        self.assertEqual(report["failed"], 1)
        self.assertIn("max_file_bytes", report["errors"]["other.docx"])


if __name__ == "__main__":
    unittest.main()
